"""Llena wms_upc_catalog con los UPC corregidos de SPEKTRUM (los 51 del Excel).

Fuente ÚNICA de verdad: la columna `UPC SHOULD BE` del Excel de correcciones,
keyed por SKU FINAL (style_final + color + size). El UPC identifica el SKU
(customer+style+color+size); description/manufacturer/país/composición son
metadata que se toman de una caja de muestra al CREAR.

Por cada SKU con UPC corregido:
  - 1 entrada en catálogo -> UPDATE su upc (si difiere)
  - 0 entradas           -> CREATE una entrada canónica
  - >1 entradas          -> FLAG (no toca; se revisa a mano: puede haber duplicados)

Solo escribe wms_upc_catalog. dry-run por defecto; APPLY=1 respalda y aplica.
"""
import os
import secrets
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from dotenv import load_dotenv

ROOT = Path(__file__).parent
load_dotenv(ROOT / ".env")
import pymongo  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

XLSX = os.environ.get("SPEKTRUM_XLSX", r"C:\Users\gerar\Downloads\Untitled spreadsheet.xlsx")
CUSTOMER = "SPEKTRUM"
BRAND = "CLCA"
APPLY = os.environ.get("APPLY") == "1"
TS = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")


def cl(v):
    return "" if v is None else str(v).strip()


def cu(v):
    s = "" if v is None else str(v).strip()
    return s[:-2] if s.endswith(".0") else s


def make_sku(style, color, size):
    import re
    parts = [re.sub(r"\s+", "-", style.upper()),
             re.sub(r"\s+", "-", color.upper())[:10] if color else "", size.upper()]
    return "-".join(p for p in parts if p)


def upc_final():
    """(style_final,color,size) -> upc corregido. Aborta si un SKU tiene 2 upc distintos."""
    ws = openpyxl.load_workbook(XLSX, data_only=True)["Sheet1"]
    H = [ws.cell(row=1, column=c).value for c in range(1, 13)]
    acc = defaultdict(set)
    for r in range(2, ws.max_row + 1):
        rec = {H[c]: ws.cell(row=r, column=c + 1).value for c in range(12)}
        if not any(rec[k] not in (None, "") for k in H):
            continue
        upc = cu(rec["UPC SHOULD BE"])
        if not upc:
            continue
        fs = cl(rec["Style SHOULD BE"]) or cl(rec["Style WAS"])
        acc[(fs, cl(rec["Color"]), cl(rec["Size"]))].add(upc)
    conf = {k: v for k, v in acc.items() if len(v) > 1}
    if conf:
        print(f"[!] ABORTA: UPC contradictorio en el Excel: {conf}")
        sys.exit(1)
    return {k: next(iter(v)) for k, v in acc.items()}


def main():
    print(f"=== fix_spektrum_catalog_upc  [{'APPLY' if APPLY else 'DRY-RUN'}]  {TS} ===")
    upcs = upc_final()
    print(f"SKUs con UPC corregido en el Excel: {len(upcs)}")

    db = pymongo.MongoClient(os.environ["MONGODB_URL"], serverSelectionTimeoutMS=20000)[
        os.environ.get("DB_NAME", "mos-system")]

    to_update, to_create, flagged = [], [], []
    for (st, c, z), upc in sorted(upcs.items()):
        entries = list(db.wms_upc_catalog.find(
            {"customer": CUSTOMER, "style": st, "color": c, "size": z}, {"_id": 1, "upc": 1}))
        if len(entries) == 1:
            if cu(entries[0].get("upc")) != upc:
                to_update.append((entries[0]["_id"], st, c, z, cu(entries[0].get("upc")), upc))
        elif len(entries) == 0:
            to_create.append((st, c, z, upc))
        else:
            flagged.append((st, c, z, upc, len(entries)))

    print(f"\n-- Plan (solo wms_upc_catalog) --")
    print(f"  UPDATE (1 entrada, upc distinto): {len(to_update)}")
    print(f"  CREATE (0 entradas): {len(to_create)}")
    print(f"  FLAG (>1 entrada, revisar a mano): {len(flagged)}")
    for st, c, z, upc, n in flagged:
        print(f"     FLAG {st}/{c}/{z} -> {upc}  ({n} entradas)")
    print("  muestra CREATE:")
    for st, c, z, upc in to_create[:8]:
        print(f"     + {st}/{c}/{z} -> {upc}")
    print("  muestra UPDATE:")
    for _id, st, c, z, old, upc in to_update[:8]:
        print(f"     ~ {st}/{c}/{z}: {old or '(vacío)'} -> {upc}")

    if not APPLY:
        print(f"\n[DRY-RUN] Nada escrito. Respaldo que se crearía: wms_upc_catalog_bak_spektrum_{TS}.")
        return

    print("\n=== RESPALDO ===")
    db.wms_upc_catalog.aggregate([{"$match": {"customer": CUSTOMER}},
                                  {"$out": f"wms_upc_catalog_bak_spektrum_{TS}"}])
    print(f"  OK: wms_upc_catalog_bak_spektrum_{TS}")

    now = datetime.now(timezone.utc).isoformat()
    for _id, st, c, z, old, upc in to_update:
        db.wms_upc_catalog.update_one({"_id": _id}, {"$set": {"upc": upc, "updated_at": now,
                                                              "updated_by": "fix_spektrum_catalog_upc"}})
    created = 0
    for st, c, z, upc in to_create:
        box = db.wms_boxes.find_one({"customer": CUSTOMER, "style": st, "color": c, "size": z},
                                    {"_id": 0, "description": 1, "manufacturer": 1,
                                     "country_of_origin": 1, "coo": 1, "fabric_content": 1}) or {}
        db.wms_upc_catalog.insert_one({
            "catalog_id": f"upc_{secrets.token_hex(6)}",
            "upc": upc, "customer": CUSTOMER, "brand": BRAND,
            "manufacturer": cl(box.get("manufacturer")) or "CL+CA",
            "style": st, "color": c, "size": z,
            "description": cl(box.get("description")),
            "country_of_origin": cl(box.get("country_of_origin")) or cl(box.get("coo")),
            "fabric_content": cl(box.get("fabric_content")),
            "sku": make_sku(st, c, z),
            "created_at": now, "created_by_name": "fix_spektrum_catalog_upc",
            "upc_source": XLSX.split("\\")[-1],
        })
        created += 1
    print(f"\n[OK] catálogo: {len(to_update)} actualizados, {created} creados, {len(flagged)} marcados sin tocar.")


if __name__ == "__main__":
    main()
