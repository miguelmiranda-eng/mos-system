"""Importa la hoja de Google "shipping miranda" al programador de envíos
(routers/scheduled_shipments.py: shipping_exports + scheduled_shipments).

Se usó el 2026-09-25 para arrancar el módulo con los envíos reales (29 exports,
432 líneas, semanas 31 AGO – 02 OCT). Sirve para volver a sincronizar mientras
la hoja siga en uso.

CÓMO LEE LA HOJA
────────────────
Una pestaña por semana ("21 SEP - 25 SEP"; TEMPLE se ignora). La fecha de cada
bloque sale del lunes de la pestaña + el nombre del día (LUNES…VIERNES). El
renglón "EXPORT#76 | PL… | TEC… | VERDE" (o un "TEC-361" suelto) precede al
renglón del día y se vuelve el encabezado de ese export. Las columnas se mapean
por el título del renglón CUSTOMER (cambia entre pestañas); status y prioridad se
reconocen por su valor. SHIP DATE, "Planning comments" y cualquier celda extra
van a las notas del renglón. Órdenes que no existen en el CRM → línea manual.

Los totales de la hoja NO se importan (el módulo los calcula): varios tienen
fórmulas con rango equivocado, y el reporte de simulación los lista.

SEGURIDAD
─────────
Todo lo importado lleva import_source=TAG. Re-ejecutar borra SOLO lo importado
antes y lo vuelve a insertar; nunca toca exports/líneas capturados en el módulo.
--undo quita la importación completa.

USO (desde la raíz del repo; usa backend/.env)
───
    python backend/scripts/import_sheet_shipping_miranda.py "shipping miranda.xlsx"           # simulación
    python backend/scripts/import_sheet_shipping_miranda.py "shipping miranda.xlsx" --apply
    python backend/scripts/import_sheet_shipping_miranda.py --undo
"""
import os, re, sys, uuid
from datetime import date, datetime, timedelta, timezone
from collections import Counter, defaultdict

import openpyxl

TAG = "sheet:shipping-miranda"
YEAR = 2026
MONTHS = {"ENE": 1, "FEB": 2, "MAR": 3, "ABR": 4, "MAY": 5, "JUN": 6, "JUL": 7, "AGO": 8,
          "SEP": 9, "OCT": 10, "NOV": 11, "DIC": 12}
DAYS = {"LUNES": 0, "MARTES": 1, "MIERCOLES": 2, "MIÉRCOLES": 2, "JUEVES": 3, "VIERNES": 4,
        "SABADO": 5, "SÁBADO": 5, "DOMINGO": 6}
STATUSES = ["READY TO SHIP", "IN SETUP", "SURTIDO A PISO", "NECK READY", "PRINTED", "PACKAGED READY",
            "QC READY", "CANCELLED", "SE MUEVE FECHA", "PRINTING", "PRIORITY"]
# El STATUS del módulo es automático desde MOS; de la hoja sólo se conserva lo
# que MOS no puede calcular. El resto se deja en None (= AUTO).
KEEP_MANUAL = {"READY TO SHIP", "PRIORITY", "SE MUEVE FECHA", "CANCELLED"}
PRIO = {"1RA": 1, "2DA": 2, "3RA": 3, "4TA": 4}


def s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, datetime):
        return v.date().isoformat()
    return re.sub(r"\s+", " ", str(v)).strip()


def num(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    t = str(v).replace(",", "").strip()
    try:
        return int(round(float(t)))
    except ValueError:
        return None


def to24(txt):
    m = re.search(r"(\d{1,2}):(\d{2})\s*(AM|PM)", txt or "", re.I)
    if not m:
        return None
    h, mi, ap = int(m.group(1)), int(m.group(2)), m.group(3).upper()
    if ap == "PM" and h != 12:
        h += 12
    if ap == "AM" and h == 12:
        h = 0
    return f"{h:02d}:{mi:02d}"


def week_start(name):
    m = re.match(r"\s*(\d{1,2})\s+([A-Z]{3})", name.upper())
    return date(YEAR, MONTHS[m.group(2)], int(m.group(1)))


def norm_place(v):
    t = s(v).upper()
    return "ST ANDREWS" if re.match(r"ST\.?\s*ANDREWS", t) else (t or None)


def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    blocks, warnings = [], []
    for name in wb.sheetnames:
        if name.strip().upper().startswith("TEMPLE"):
            continue
        ws = wb[name]
        monday = week_start(name)
        cols = {}
        pending = {}           # encabezado de export que precede al renglón del día
        cur = None
        for ri, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = list(row)
            txt = [s(c) for c in cells]
            a = txt[0] if txt else ""
            au = a.upper()
            nonempty = [t for t in txt if t]
            if not nonempty:
                continue
            if au.startswith("CUSTOMER"):
                cols = {}
                for i, h in enumerate(txt):
                    hu = h.upper()
                    if hu:
                        cols.setdefault(hu, i)
                continue
            if au.startswith("EXPORT#"):
                pending = {"export_no": num(re.sub(r"\D", "", a)) or None}
                for t in txt[1:]:
                    tu = t.upper()
                    if not t:
                        continue
                    if tu.startswith("PL") and "pl" not in pending:
                        pending["pl"] = t
                    elif re.search(r"\b(VERDE|ROJO)\b", tu):
                        pending["light"] = "VERDE" if "VERDE" in tu else "ROJO"
                    elif "truck" not in pending:
                        pending["truck"] = t
                continue
            if re.match(r"^(TEC|RABON)", au) and len(nonempty) == 1:
                pending.setdefault("truck", a)
                continue
            dm = re.match(r"^(LUNES|MARTES|MI[EÉ]RCOLES|JUEVES|VIERNES|S[AÁ]BADO|DOMINGO)", au)
            if dm:
                day = monday + timedelta(days=DAYS[dm.group(1)])
                joined = " ".join(txt)
                cur = {"sheet": name.strip(), "row": ri, "date": day.isoformat(),
                       "cutoff": to24(joined.split("EXPORT HR")[0]) if "CORTE" in joined.upper() else None,
                       "export_time": to24(joined.split("EXPORT HR")[1]) if "EXPORT HR" in joined.upper() else None,
                       "export_no": pending.get("export_no"), "pl": pending.get("pl"),
                       "truck": pending.get("truck"), "light": pending.get("light"),
                       "notes": [], "lines": [], "sheet_total": None}
                pending = {}
                blocks.append(cur)
                continue
            if cur is None:
                warnings.append(f"{name.strip()} r{ri}: fuera de bloque, ignorado: {nonempty[:8]}")
                continue
            ci = cols.get("ORDER", 3)
            order_raw = txt[ci] if ci < len(txt) else ""
            if not a and not order_raw:
                pcs_i = cols.get("PCS", 6)
                tot = num(cells[pcs_i]) if pcs_i < len(cells) else None
                if tot is not None and len(nonempty) == 1:
                    cur["sheet_total"] = tot
                else:
                    warnings.append(f"{name.strip()} r{ri}: renglón suelto: {nonempty[:8]}")
                continue
            m = re.match(r"^\s*(\d+)", order_raw)
            if not a or not m:
                cur["notes"].append(" ".join(nonempty))
                continue
            get = lambda key, default=None: (txt[cols[key]] if key in cols and cols[key] < len(txt) else (txt[default] if default is not None and default < len(txt) else ""))
            pcs_i = cols.get("PCS", 6)
            status_raw = get("STATUS", pcs_i + 1).upper()
            extra, status, prio = [], None, None
            # Celdas a la derecha de PCS: status, prioridad, notas, etc.
            for i in range(pcs_i + 1, len(txt)):
                t = txt[i]
                if not t:
                    continue
                tu = t.upper()
                pm = re.match(r"^(1RA|2DA|3RA|4TA)\s+PRIORIDAD", tu)
                if tu in STATUSES and status is None:
                    status = tu
                elif pm and prio is None:
                    prio = PRIO[pm.group(1)]
                elif i in (cols.get("SHIPPING FROM"), cols.get("CARRIER")):
                    continue
                elif i == cols.get("SHIP DATE"):
                    extra.append(f"SHIP DATE {t}")
                else:
                    extra.append(t)
            if status is None and status_raw and status_raw not in STATUSES:
                pass  # ya quedó en extra
            cur["lines"].append({
                "row": ri,
                "order_number": m.group(1),
                "client": get("CUSTOMER", 0) or None,
                "shipping_no": get("SHIPPING#", 1) or None,
                "delivery_to": norm_place(get("DELIVER TO")) if "DELIVER TO" in cols else None,
                "branding": get("BRANDING") or None,
                "customer_po": get("CUSTOMER PO.") or None,
                "design_num": get("DESIGN #") or None,
                "pcs": num(cells[pcs_i]) if pcs_i < len(cells) else None,
                "status": status,
                "priority": prio,
                "ship_notes": " · ".join(extra) or None,
                "ship_from": norm_place(get("SHIPPING FROM")) if "SHIPPING FROM" in cols else None,
                "carrier": (get("CARRIER") or None) if "CARRIER" in cols else None,
            })
    return blocks, warnings


def _db():
    """Base de datos de backend/.env (misma que usa el servidor)."""
    from dotenv import load_dotenv
    import pymongo
    load_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".env"))
    url = os.environ.get("MONGODB_URL") or os.environ.get("MONGO_URL")
    if not url:
        sys.exit("Falta MONGODB_URL en backend/.env")
    return pymongo.MongoClient(url, serverSelectionTimeoutMS=20000)[os.environ.get("DB_NAME", "mos-system")]


def undo():
    db = _db()
    ids = [e["export_id"] for e in db.shipping_exports.find({"import_source": TAG}, {"export_id": 1})]
    nl = db.scheduled_shipments.delete_many({"import_source": TAG}).deleted_count
    ne = db.shipping_exports.delete_many({"import_source": TAG}).deleted_count
    # Líneas capturadas a mano DENTRO de un export importado quedarían huérfanas: se avisa.
    orphan = db.scheduled_shipments.count_documents({"export_id": {"$in": ids}}) if ids else 0
    print(f"Quitado: {ne} exports y {nl} líneas importadas."
          + (f" OJO: {orphan} líneas capturadas a mano apuntan a exports borrados." if orphan else ""))


def main():
    if "--undo" in sys.argv:
        return undo()
    if len(sys.argv) < 2 or sys.argv[1].startswith("--"):
        sys.exit(__doc__)
    path = sys.argv[1]
    apply = "--apply" in sys.argv
    blocks, warnings = parse(path)
    blocks = [b for b in blocks if b["lines"] or b["export_no"] or b["notes"]]
    # Cuadre contra los totales de la hoja
    mism = [(b["sheet"], b["date"], b["export_no"], sum(l["pcs"] or 0 for l in b["lines"]), b["sheet_total"])
            for b in blocks if b["sheet_total"] is not None and sum(l["pcs"] or 0 for l in b["lines"]) != b["sheet_total"]]
    nlines = sum(len(b["lines"]) for b in blocks)
    print(f"Bloques (exports) con órdenes: {len(blocks)} · líneas: {nlines}")
    per_week = Counter(b["sheet"] for b in blocks)
    for k, v in per_week.items():
        print(f"  {k}: {v} exports, {sum(len(b['lines']) for b in blocks if b['sheet']==k)} líneas, "
              f"{sum(l['pcs'] or 0 for b in blocks if b['sheet']==k for l in b['lines']):,} pzs")
    print(f"Totales que NO cuadran con la hoja: {len(mism)}")
    for x in mism:
        print("   ", x)
    print(f"Status: {Counter(l['status'] for b in blocks for l in b['lines'])}")
    print(f"Prioridad: {Counter(l['priority'] for b in blocks for l in b['lines'])}")
    print(f"Avisos del parser: {len(warnings)}")
    for w in warnings:
        print("   ", w)

    db = _db()
    nums = sorted({l["order_number"] for b in blocks for l in b["lines"]})
    found = {}
    for i in range(0, len(nums), 50):
        for o in db.orders.find({"order_number": {"$in": nums[i:i+50]}, "board": {"$ne": "PAPELERA DE RECICLAJE"}},
                                {"_id": 0, "order_number": 1, "client": 1, "cancel_date": 1}):
            found[o["order_number"]] = o
    missing = [n for n in nums if n not in found]
    print(f"Órdenes distintas: {len(nums)} · en el CRM: {len(found)} · NO en el CRM (irán como manuales): {len(missing)}")
    print("   ", missing)
    mine = db.shipping_exports.count_documents({"import_source": {"$ne": TAG}})
    prev = db.shipping_exports.count_documents({"import_source": TAG})
    print(f"Exports capturados en el módulo (no se tocan): {mine} · de una importación previa (se reemplazan): {prev}")
    if not apply:
        print("\nSIMULACIÓN: no se escribió nada. Usa --apply para importar.")
        return

    old_ids = [e["export_id"] for e in db.shipping_exports.find({"import_source": TAG}, {"export_id": 1})]
    if old_ids:
        db.scheduled_shipments.delete_many({"export_id": {"$in": old_ids}, "import_source": TAG})
        db.shipping_exports.delete_many({"import_source": TAG})
    now = datetime.now(timezone.utc).isoformat()
    who = {"created_by": "import", "created_by_name": "Importado de hoja shipping miranda"}
    pos_by_day = defaultdict(int)
    for d in db.shipping_exports.find({"import_source": {"$ne": TAG}}, {"date": 1}):
        pos_by_day[d["date"]] += 1
    exp_docs, line_docs = [], []
    for b in blocks:
        eid = str(uuid.uuid4())
        exp_docs.append({
            "export_id": eid, "date": b["date"], "position": pos_by_day[b["date"]],
            "export_no": b["export_no"], "pl_numbers": b["pl"], "truck": b["truck"],
            "customs_light": b["light"], "cutoff_time": b["cutoff"] or "15:00",
            "export_time": b["export_time"] or "17:00",
            "notes": " · ".join(b["notes"]) or None,
            **who, "created_at": now, "updated_at": now, "import_source": TAG,
        })
        pos_by_day[b["date"]] += 1
        d = date.fromisoformat(b["date"])
        for i, l in enumerate(b["lines"]):
            manual = l["order_number"] not in found
            line_docs.append({
                "shipment_id": str(uuid.uuid4()), "order_number": l["order_number"], "export_id": eid,
                "ship_date": b["date"], "scheduled_export_date": b["date"],
                "scheduled_year": d.year, "scheduled_month": d.month, "position": i,
                "pcs": l["pcs"], "shipping_no": l["shipping_no"], "delivery_to": l["delivery_to"],
                "ship_from": l["ship_from"], "carrier": l["carrier"],
                "status": l["status"] if l["status"] in KEEP_MANUAL else None,
                "cancel_date_at_schedule": (found.get(l["order_number"]) or {}).get("cancel_date") or None,
                "priority": l["priority"], "ship_notes": l["ship_notes"], "manual": manual,
                "manual_fields": ({"client": l["client"], "branding": l["branding"],
                                   "customer_po": l["customer_po"], "design_num": l["design_num"]}
                                  if manual else {}),
                **who, "created_at": now, "updated_at": now, "import_source": TAG,
            })
    for i in range(0, len(exp_docs), 25):
        db.shipping_exports.insert_many(exp_docs[i:i+25])
    for i in range(0, len(line_docs), 25):
        db.scheduled_shipments.insert_many(line_docs[i:i+25])
    print(f"\nIMPORTADO: {len(exp_docs)} exports, {len(line_docs)} líneas "
          f"({sum(1 for x in line_docs if x['manual'])} manuales).")
    print("Verificación:", db.shipping_exports.count_documents({"import_source": TAG}),
          db.scheduled_shipments.count_documents({"import_source": TAG}))


if __name__ == "__main__":
    main()
