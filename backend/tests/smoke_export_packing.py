"""Smoke del motor de PACKING LIST de exportación (services/export_packing.py).

Fija: lectura del DIGITAL PACKING LIST (formato de la hoja de empaque: Box Qty,
Pallet #, TOTAL SHIPPED, resumen % × país × talla y tablas por talla), regla de
fuente (WMS si cuadra con TOTAL SHIPPED, si no DPL), CUST PO de licencias,
orden sin cliente, avisos, y el Excel resultante (formato del PL GTS).

Los DPL reales viven en Google; aquí se arma uno sintético con el mismo
formato y se sustituye la descarga (no hay red en el smoke).

SEGURIDAD: base DESECHABLE, se niega contra producción, se borra al terminar.

    set MONGODB_URL=mongodb://localhost:27017
    python backend/tests/smoke_export_packing.py
"""
import asyncio
import base64
import io
import os
import sys
import zipfile

BE = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SMOKE_DB = os.environ.get("SMOKE_DB_NAME", "mos-smoke-export-packing")
MONGO = os.environ.get("MONGODB_URL") or os.environ.get("MONGO_URL")
if not MONGO:
    sys.exit("Falta MONGODB_URL")
if SMOKE_DB == os.environ.get("PROD_DB_NAME", "mos-system"):
    sys.exit("NEGADO: base de producción")
os.environ.update({"MONGODB_URL": MONGO, "MONGO_URL": MONGO, "DB_NAME": SMOKE_DB})
for k, v in (("JWT_SECRET", "s"), ("MASTER_API_KEY", "m"), ("INTERNAL_SYNC_TOKEN", "t"),
             ("DISABLE_SCHEDULERS", "1"), ("ENV", "local")):
    os.environ.setdefault(k, v)
sys.path.insert(0, BE)
os.chdir(BE)

import openpyxl  # noqa: E402
import pymongo  # noqa: E402
from passlib.hash import bcrypt  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
raw = pymongo.MongoClient(MONGO)
sdb = raw[SMOKE_DB]
ok = fail = 0


def check(name, cond, detail=""):
    global ok, fail
    ok, fail = (ok + 1, fail) if cond else (ok, fail + 1)
    print(f"   {'PASS' if cond else 'FAIL'}  {name}" + ("" if cond else f"  {detail}"))


def fake_dpl(rows, summary, garment="SHORT SLEEVE", summary_override=None):
    """DPL con el layout real: encabezado en fila 7, TOTAL SHIPPED en 33,
    tablas por talla desde 43 y resumen en 71."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PACKING LIST"
    head = ["STORE PO", "Garment Type", "Garment Color", None, None, "XS or OSFA", "SM", "MD", "LG", "XL", "2XL",
            "3XL", "4XL", "TOTAL", "Box Qty", "Pcs per  box", "Box Dimension (LxWxH)", "Box Weight (lbs)",
            "Pallet #", "Pallet Dimension (LxWxH)"]
    for i, h in enumerate(head, start=1):
        if h:
            ws.cell(7, i, h)
    ws["M4"], ws["O4"] = "Date Packed:", "29/9/26"
    ws["M5"], ws["O5"] = "Packer Name:", "luis"
    col = {"S": 7, "M": 8, "L": 9, "XL": 10, "2X": 11}
    for i, (sz, qty, boxes, pallet) in enumerate(rows):
        r = 8 + i
        if i == 0:
            ws.cell(r, 2, garment)
        ws.cell(r, col[sz], qty)
        ws.cell(r, 15, boxes)
        if pallet:
            ws.cell(r, 19, pallet)
    ws["A33"] = "TOTAL SHIPPED:"
    # Tablas por talla (SIZE SM en F43 / SIZE MD en O43)
    ws["F43"], ws["O43"] = "SIZE SM", "SIZE MD"
    for c, lab in ((6, "Percentage"), (10, "Country of Origin"), (13, "Qty"), (15, "Percentage"),
                   (18, "Country of Origin"), (20, "Qty")):
        ws.cell(44, c, lab)
    # Resumen final
    for c, lab in ((1, "Percentage"), (3, "Country of Origin"), (5, "XS or OSFA"), (6, "SM"), (7, "MD"),
                   (8, "LG"), (9, "XL"), (10, "2XL"), (13, "TOTAL")):
        ws.cell(71, c, lab)
    r = 73
    for fab, cty, sizes in (summary_override or summary):
        ws.cell(r, 1, fab)
        ws.cell(r, 3, cty)
        for sz, q in sizes.items():
            ws.cell(r, {"S": 6, "M": 7, "L": 8, "XL": 9, "2X": 10}[sz], q)
        r += 1
    ws.cell(86, 1, "TOTAL")
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def sembrar():
    for c in ["orders", "users", "user_sessions", "scheduled_shipments", "shipping_exports", "wms_movements",
              "wms_boxes", "wms_neck_counts", "activity_logs", "shipping_movements"]:
        sdb[c].delete_many({})
    link = lambda i: [{"url": f"https://docs.google.com/spreadsheets/d/FAKE{i}/edit", "description": "DIGITAL PACKING LIST."}]
    sdb.orders.insert_many([
        {"order_id": "o1", "order_number": "3319", "client": "GOODIE TWO SLEEVES", "branding": "SPENCERS",
         "customer_po": "22726", "store_po": "325887", "store_po#": "325887", "design_#": "DNC0014M1000",
         "color": "BLACK", "board": "FINAL BILL", "sizes": {"S": 216, "M": 612}, "links": link(1)},
        {"order_id": "o2", "order_number": "3384", "client": "GOODIE TWO SLEEVES", "branding": "SPENCERS WARNER",
         "customer_po": "22767", "store_po": "P028209", "store_po#": "P028209 - 325985", "design_#": "SOD0002M1000",
         "color": "BLACK", "board": "FINAL BILL", "sizes": {"M": 100}, "links": link(2)},
        {"order_id": "o3", "order_number": "3383", "client": "", "branding": "SPENCERS WARNER",
         "customer_po": "22766", "store_po#": "P028208 - 325984", "design_#": "MYC0001M1000",
         "color": "BLACK", "board": "FINAL BILL", "sizes": {"L": 50}},
        {"order_id": "o4", "order_number": "3206", "client": "SPEKTRUM", "branding": "CULTURE KINGS",
         "customer_po": "4004627", "store_po": "4004627", "design_#": "VENOM VINTAGE TEE",
         "color": "BLACK ACIDWASH", "board": "FINAL BILL", "sizes": {"M": 30}},
    ])
    sdb.wms_boxes.insert_many([
        {"box_id": "B1", "country_of_origin": "NICARAGUA", "fabric_content": "100% COTTON", "description": "MENS SS"},
        {"box_id": "B2", "country_of_origin": "REPUBLICA DOMINICANA", "fabric_content": "100% COTTON", "description": "MENS SS"},
        {"box_id": "B3", "country_of_origin": "HAITI", "fabric_content": "100% COTTON", "description": "MENS SS"},
    ])
    mv = lambda num, sz, box, q: {"type": "pick_deduction", "details": {"order_number": num, "size": sz,
                                                                         "boxes": [{"box_id": box, "taken": q}]}}
    sdb.wms_movements.insert_many([
        mv("3319", "S", "B1", 216), mv("3319", "M", "B1", 612),
        mv("3384", "M", "B2", 100),
        mv("3383", "L", "B3", 50),
        mv("3206", "M", "B3", 30),
    ])
    sdb.users.insert_one({"user_id": "u", "email": "u@test.local", "name": "Envíos", "password_hash": bcrypt.hash("p"),
                          "role": "supersu", "admin_level": 5, "active": True})


async def main():
    sembrar()
    from services import export_packing as ep

    print("== Lectura del DPL ==")
    dpl_3319 = fake_dpl([("S", 216, 6, 1), ("M", 612, 17, 2)], [("100% Cotton", "Nicaragua", {"S": 216, "M": 612})])
    p = ep.parse_dpl(dpl_3319)
    check("cajas = suma de Box Qty", p["boxes"] == 23, p["boxes"])
    check("tarimas = Pallet # distintos", p["pallets"] == 2, p["pallets"])
    check("TOTAL SHIPPED por talla", p["by_size"] == {"S": 216, "M": 612}, p["by_size"])
    check("resumen país × talla", p["rows"] == [{"fabric": "100% COTTON", "country": "NICARAGUA", "sizes": {"S": 216, "M": 612}}], p["rows"])
    check("Garment Type, fecha y empacador", p["garment_type"] == "SHORT SLEEVE" and p["packer"] == "luis", p)
    check("país normalizado (Rep. Dominicana → REP DOMINICANA)", ep.norm_country("Rep. Dominicana") == "REP DOMINICANA")
    check("descripción del WMS (MENS SS → SHORT SLEEVE)", ep.norm_description("MENS SS") == "SHORT SLEEVE")
    check("CUST PO de licencia toma el PO de la tienda", ep.cust_po({"store_po#": "P028209 - 325985"}) == 325985)

    # DPLs falsos por enlace (sin red)
    dpl_3384 = fake_dpl([("M", 100, 3, 1)], [("100% Cotton", "Nicaragua", {"M": 100})])   # país distinto al WMS

    async def fake_fetch(client, url):
        data = {"FAKE1": dpl_3319, "FAKE2": dpl_3384}.get(ep.DPL_RE.search(url).group(1))
        return (ep.parse_dpl(data), None) if data else (None, "no existe")
    ep.fetch_dpl = fake_fetch

    from httpx import ASGITransport, AsyncClient
    from server import app
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://smoke") as c:
        await c.post("/api/auth/login", json={"email": "u@test.local", "password": "p"})
        API = "/api/scheduled-shipments"
        exp = (await c.post(f"{API}/exports", json={"date": "2026-09-29"})).json()
        await c.put(f"{API}/exports/{exp['export_id']}", json={
            "export_no": 85, "pl_numbers": "PLGTS 09-26-0085", "truck": "TEC.361 / 53144",
            "transport_company": "TECMA TRANSPORTATION", "driver_name": "JAIME", "license_plate": "BU9842",
            "seal_numbers": "014093 // 014615"})
        add = (await c.post(f"{API}/lines", json={"export_id": exp["export_id"], "order_numbers": "3319 3384 3383"})).json()["added"]
        for ln in add:
            await c.put(f"{API}/{ln['shipment_id']}", json={"delivery_to": "ST ANDREWS"})

        print("\n== PL de un envío (export) con una orden ==")
        uno = (await c.post(f"{API}/exports", json={"date": "2026-09-30"})).json()
        await c.put(f"{API}/exports/{uno['export_id']}", json={
            "export_no": 85, "pl_numbers": "PLGTS 09-26-0085", "truck": "TEC.361 / 53144",
            "transport_company": "TECMA TRANSPORTATION", "driver_name": "JAIME", "license_plate": "BU9842",
            "seal_numbers": "014093 // 014615"})
        l1 = (await c.post(f"{API}/lines", json={"export_id": uno["export_id"], "order_numbers": "3319"})).json()["added"][0]
        await c.put(f"{API}/{l1['shipment_id']}", json={"delivery_to": "ST ANDREWS"})
        r = await c.post(f"{API}/exports/{uno['export_id']}/packing")
        d = r.json()
        check("responde archivo xlsx con nombre del PL", r.status_code == 200 and d["filename"].startswith("PL GTS 09-26-0085")
              and d["filename"].endswith(".xlsx"), d.get("filename"))
        ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(d["content_b64"])))["PACKING"]
        check("encabezado: # PACKING LIST y SHIP TO", ws["V1"].value == "PLGTS 09-26-0085" and ws["V6"].value == "TSC BROKER")
        check("CUSTOMER y SHIP TO como los escribe Envíos", ws["U11"].value == "SPENCER" and ws["V11"].value == "ST ANDREW",
              (ws["U11"].value, ws["V11"].value))
        check("no existe el PL por renglón (es por envío)",
              (await c.post(f"{API}/lines/{l1['shipment_id']}/packing")).status_code in (404, 405))
        fila = [ws.cell(11, i).value for i in range(1, 9)]
        check("datos de la orden (INV/PO/CUST PO/STYLE/COLOR/CONTENT/DESC/ORIGEN)",
              fila == [3319, 22726, 325887, "DNC0014M1000", "BLACK", "100% COTTON", "SHORT SLEEVE", "NICARAGUA"], fila)
        check("tallas SM/MD en su columna", ws["J11"].value == 216 and ws["K11"].value == 612)
        check("cajas y tarimas del DPL", ws["S11"].value == 23 and ws["T11"].value == 2, (ws["S11"].value, ws["T11"].value))
        check("pie: transporte, chofer, placas, camión, sello y No.Eco",
              [ws[f"E{r}"].value for r in range(14, 21)] == ["TECMA TRANSPORTATION", "JAIME", "BU9842", "TEC.361", None,
                                                            "014093 // 014615", 53144],
              [ws[f"E{r}"].value for r in range(14, 21)])
        check("fuente WMS (cuadra con TOTAL SHIPPED)", d["summary"][0]["source"].startswith("WMS"), d["summary"])
        check("sin avisos cuando todo cuadra", d["warnings"] == [], d["warnings"])

        print("\n== PL del export completo ==")
        r = await c.post(f"{API}/exports/{exp['export_id']}/packing")
        d = r.json()
        names = [d["filename"]]
        if d["filename"].endswith(".zip"):
            names = zipfile.ZipFile(io.BytesIO(base64.b64decode(d["content_b64"]))).namelist()
        check("un solo PL GTS (la orden sin cliente se suma al cliente del export)", len(names) == 1 and "GTS" in names[0], names)
        ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(d["content_b64"])))["PACKING"]
        invs = [ws.cell(r, 1).value for r in range(11, 14)]
        check("las 3 órdenes en el PL", invs == [3319, 3384, 3383], invs)
        check("CUST PO de licencia = PO de tienda", ws["C12"].value == 325985, ws["C12"].value)
        w = " | ".join(d["warnings"])
        check("aviso: país distinto entre DPL y WMS", "#3384: país de origen distinto entre DPL y WMS" in w, w)
        check("aviso: orden sin cliente", "#3383: la orden no tiene cliente" in w, w)
        check("aviso: orden sin DPL (faltan cajas/tarimas)", "#3383: sin enlace al DIGITAL PACKING LIST" in w, w)
        check("totales con fórmula", str(ws["N18"].value or "").startswith("=SUM(R11:R13)") or
              any(str(ws.cell(r, 14).value or "").startswith("=SUM(R11:R13)") for r in range(14, 25)))

        print("\n== Envío con varios clientes: un packing por cliente ==")
        await c.post(f"{API}/lines", json={"export_id": exp["export_id"], "order_numbers": "3206"})
        r = await c.get(f"{API}/exports/{exp['export_id']}/packing/clients")
        cl = {x["code"]: x for x in r.json()["clients"]}
        check("lista de clientes del envío (GTS 3 órdenes con la sin cliente, SKT 1)",
              set(cl) == {"GTS", "SKT"} and cl["GTS"]["count"] == 3 and cl["SKT"]["count"] == 1, r.json())
        check("PL de cada cliente (SKT sale del EXPORT# si no está capturado)",
              cl["GTS"]["pl_number"] == "PLGTS 09-26-0085" and cl["SKT"]["pl_number"] == "PLSKT 09-26-0085", cl)
        r = await c.post(f"{API}/exports/{exp['export_id']}/packing", params={"client": "skt"})
        d = r.json()
        ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(d["content_b64"])))["PACKING"]
        check("packing de un solo cliente: xlsx con su PL y sólo sus órdenes",
              d["filename"].startswith("PL SKT 09-26-0085") and ws["V1"].value == "PLSKT 09-26-0085"
              and ws["A11"].value == 3206 and not ws["A12"].value, (d["filename"], ws["V1"].value, ws["A11"].value, ws["A12"].value))
        check("los avisos son sólo de ese cliente", not any(w.startswith(("#3383", "#3384")) for w in d["warnings"]), d["warnings"])
        r = await c.post(f"{API}/exports/{exp['export_id']}/packing")
        names = zipfile.ZipFile(io.BytesIO(base64.b64decode(r.json()["content_b64"]))).namelist()
        check("sin cliente: todos los packings del envío en un .zip", sorted(n[:6] for n in names) == ["PL GTS", "PL SKT"], names)
        r = await c.post(f"{API}/exports/{exp['export_id']}/packing", params={"client": "ABC"})
        check("cliente que no va en el envío → 404", r.status_code == 404, r.status_code)

        print("\n== DPL con resumen mal capturado ==")
        mala = fake_dpl([("S", 216, 6, 1), ("M", 612, 17, 1)], [("100% Cotton", "Nicaragua", {"S": 216, "M": 600})])
        pm = ep.parse_dpl(mala)
        check("resumen que no cuadra con TOTAL SHIPPED no se usa como tal", ep.size_totals(pm["rows_summary"]) != pm["by_size"])
        r = await c.post(f"{API}/exports/xxx/packing")
        check("export inexistente → 404", r.status_code == 404, r.status_code)


try:
    asyncio.run(main())
finally:
    raw.drop_database(SMOKE_DB)
    print(f"\n{ok} PASS · {fail} FAIL")
    sys.exit(1 if fail else 0)
